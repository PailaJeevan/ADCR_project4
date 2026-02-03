# File Reader Module


import pdfplumber # PDF Reading Library
import docx # DOCX Reading Library
import openpyxl # XLSX Reading Library


class FileReading: # File Reader Class for Multiple Formats

    def reading_file(self, file_path):
        try:
            
            if file_path.endswith('.pdf'): # Read PDF files
                return self.reading_pdf(file_path)
            elif file_path.endswith('.docx'):# Read DOCX files
                return self.reading_docx(file_path)
            elif file_path.endswith('.txt'):# Read TXT files
                return self.reading_txt(file_path)
            elif file_path.endswith('.xlsx'):# Read XLSX files
                return self.reading_xlsx(file_path)
            else:
                raise ValueError("Unsupported file format")
        except Exception as e:
            print(f"Error reading file {file_path}: {e}")
            return None

    def reading_pdf(self, file_path): # Read and extract text from PDF files
        text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text

    def reading_docx(self, file_path): # Read and extract text from DOCX files
        doc = docx.Document(file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text

    def reading_txt(self, file_path): # Read and extract text from TXT files
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def reading_xlsx(self, file_path): # Read and extract text from XLSX files
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        text = ""

        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    text += str(cell) + " "
            text += "\n"

        return text
